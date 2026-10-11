from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from itertools import combinations
import json
from pathlib import Path
from statistics import mean
import sys
from typing import Any, Literal, Mapping, Sequence

import numpy as np
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from workflow.common.agent_io import read_ie_paper
from evaluation.sbms import (
    QwenEmbeddingClient,
    cosine_similarity_matrix,
    resolve_embedding_config,
    sbms_from_similarity,
)


ComparisonMode = Literal["label_only", "label_plus_evidence"]
EXPERTS = ("Wenjing Ma", "Jiaying Lu", "Jasmine Zhou", "Junyuan Zhang")
MODES: tuple[ComparisonMode, ...] = ("label_only", "label_plus_evidence")
PAPERS = (
    "Chu_NatMed_2023",
    "Ovarian_Nature_2022",
    "BreastAtlas_NatGenet_2021",
    "CARTMultiomics_bioRxiv_2024",
    "OralGB_CancerSci_2023",
    "FL_BloodCancerDiscov_2022",
    "TNBCStroma_EMBOJ_2020",
)
CONDITIONS = (
    "deepseek_v4_flash_thinking",
    "deepseek_v4_flash_non_thinking",
    "deepseek_v4_pro_thinking",
    "deepseek_v4_pro_non_thinking",
    "mimo_v2_5_pro_openrouter_thinking",
    "gpt5_5_openrouter_medium",
)
GEMINI_CONDITION = "gemini_3_5_flash_openrouter_thinking_high"
CONDITION_LABELS = {
    "deepseek_v4_flash_thinking": "DeepSeek V4 Flash thinking (max)",
    "deepseek_v4_flash_non_thinking": "DeepSeek V4 Flash non-thinking",
    "deepseek_v4_pro_thinking": "DeepSeek V4 Pro thinking (max)",
    "deepseek_v4_pro_non_thinking": "DeepSeek V4 Pro non-thinking",
    "mimo_v2_5_pro_openrouter_thinking": "MiMo V2.5 Pro OpenRouter thinking (high)",
    "gpt5_5_openrouter_medium": "GPT-5.5 OpenRouter thinking (medium)",
    GEMINI_CONDITION: "Gemini 3.5 Flash OpenRouter thinking (high)",
}
EXPECTED_HUMAN_COUNTS = {
    "Chu_NatMed_2023": (41, 41, 41, 42),
    "Ovarian_Nature_2022": (13, 13, 13, 15),
    "BreastAtlas_NatGenet_2021": (10, 10, 10, 11),
    "CARTMultiomics_bioRxiv_2024": (7, 7, 7, 8),
    "OralGB_CancerSci_2023": (7, 9, 8, 7),
    "FL_BloodCancerDiscov_2022": (7, 7, 7, 7),
    "TNBCStroma_EMBOJ_2020": (7, 10, 7, 7),
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def mean_exact(values: Sequence[float], *, expected: int) -> float:
    if len(values) != expected:
        raise ValueError(f"Expected {expected} values, found {len(values)}")
    return float(mean(values))


def evidence_text(value: Any) -> str:
    if isinstance(value, str):
        return " ".join(value.split())
    if isinstance(value, (list, tuple)):
        return " ".join(" ".join(str(item).split()) for item in value if str(item).strip())
    return " ".join(str(value).split()) if value is not None else ""


def build_term_texts(
    records: Sequence[Mapping[str, Any]],
    *,
    mode: ComparisonMode,
) -> list[str]:
    texts: list[str] = []
    for record in records:
        label = str(record.get("label", "")).strip()
        if not label:
            continue
        if mode == "label_only":
            texts.append(label)
            continue
        evidence = evidence_text(record.get("evidence"))
        texts.append(f"{label}\nEvidence: {evidence}" if evidence else label)
    return texts


def load_human_standard(path: Path) -> dict[str, dict[str, list[dict[str, str]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if set(workbook.sheetnames) != set(PAPERS) or len(workbook.sheetnames) != len(PAPERS):
        raise ValueError(f"Unexpected human workbook sheets: {workbook.sheetnames}")
    papers: dict[str, dict[str, list[dict[str, str]]]] = {}
    for paper in PAPERS:
        rows = list(workbook[paper].iter_rows(values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        experts: dict[str, list[dict[str, str]]] = {}
        for expert in EXPERTS:
            label_column = headers.index(f"{expert} (subtype)")
            evidence_column = headers.index(f"{expert} (evidence)")
            records: list[dict[str, str]] = []
            for row in rows[1:]:
                label_value = row[label_column] if label_column < len(row) else None
                evidence_value = row[evidence_column] if evidence_column < len(row) else None
                label = str(label_value or "").strip()
                if not label:
                    continue
                records.append(
                    {
                        "label": label,
                        "evidence": str(evidence_value or "").strip(),
                    }
                )
            experts[expert] = records
        counts = tuple(len(experts[expert]) for expert in EXPERTS)
        if counts != EXPECTED_HUMAN_COUNTS[paper]:
            raise ValueError(f"Unexpected human counts for {paper}: {counts}")
        papers[paper] = experts
    workbook.close()
    return papers


def resolve_output_path(run_root: Path, condition: str, paper: str) -> Path:
    directory = run_root / "ie_agent" / condition
    matches = sorted(directory.glob(f"{paper}_prompt_V2_{condition}_direct.json"))
    if len(matches) != 1:
        raise ValueError(f"Expected one output for {condition}/{paper}, found {matches}")
    return matches[0]


def resolve_condition_paths(
    run_root: Path,
    *,
    gemini_run_root: Path | None = None,
) -> dict[str, dict[str, Path]]:
    paths = {
        condition: {paper: resolve_output_path(run_root, condition, paper) for paper in PAPERS}
        for condition in CONDITIONS
    }
    if gemini_run_root is not None:
        paths[GEMINI_CONDITION] = {
            paper: resolve_output_path(gemini_run_root, GEMINI_CONDITION, paper)
            for paper in PAPERS
        }
    return paths


def resolve_result_paths(output_dir: Path, *, include_gemini: bool) -> tuple[Path, Path]:
    suffix = "_with_gemini35_flash_high" if include_gemini else ""
    return (
        output_dir / f"ie_sbms{suffix}.json",
        output_dir / f"ie_sbms{suffix}.csv",
    )


def load_model_records(path: Path) -> list[dict[str, Any]]:
    paper = read_ie_paper(path)
    return [{"label": item.original_label, "evidence": item.evidence} for item in paper.items]


def load_condition_models(
    paths: Mapping[str, Mapping[str, Path]],
) -> dict[str, dict[str, list[dict[str, Any]]]]:
    return {
        condition: {paper: load_model_records(path) for paper, path in paper_paths.items()}
        for condition, paper_paths in paths.items()
    }


class SharedEmbeddingCache:
    def __init__(self, config_path: Path) -> None:
        self.config = resolve_embedding_config(config_path=config_path)
        self.client = QwenEmbeddingClient(self.config)
        self.values: dict[str, np.ndarray] = {}
        self.dimension: int | None = None

    def prepare(self, texts: Sequence[str]) -> None:
        unique = list(dict.fromkeys(texts))
        if not unique:
            raise ValueError("No texts to embed")
        embeddings = self.client.embed_terms(unique, batch_size=10)
        if embeddings.shape[0] != len(unique) or embeddings.ndim != 2:
            raise ValueError(f"Unexpected embedding shape: {embeddings.shape}")
        if not np.isfinite(embeddings).all():
            raise ValueError("Embedding response contains non-finite values")
        self.dimension = int(embeddings.shape[1])
        self.values = {text: embeddings[index] for index, text in enumerate(unique)}

    def matrix(self, texts: Sequence[str]) -> np.ndarray:
        return np.vstack([self.values[text] for text in texts])


def score_terms(gold: Sequence[str], pred: Sequence[str], cache: SharedEmbeddingCache) -> float:
    matrix = cosine_similarity_matrix(cache.matrix(gold), cache.matrix(pred))
    score, _, _, _, _ = sbms_from_similarity(matrix)
    return float(score)


def collect_texts(
    human: Mapping[str, Mapping[str, Sequence[Mapping[str, Any]]]],
    models: Mapping[str, Mapping[str, Sequence[Mapping[str, Any]]]],
    *,
    conditions: Sequence[str] = CONDITIONS,
) -> list[str]:
    texts: list[str] = []
    for mode in MODES:
        for paper in PAPERS:
            for expert in EXPERTS:
                texts.extend(build_term_texts(human[paper][expert], mode=mode))
            for condition in conditions:
                texts.extend(build_term_texts(models[condition][paper], mode=mode))
    return texts


def evaluate_condition(
    condition: str,
    human: Mapping[str, Mapping[str, Sequence[Mapping[str, Any]]]],
    models: Mapping[str, Mapping[str, Sequence[Mapping[str, Any]]]],
    cache: SharedEmbeddingCache,
) -> dict[str, Any]:
    modes: dict[str, Any] = {}
    for mode in MODES:
        paper_results: dict[str, Any] = {}
        paper_scores: list[float] = []
        for paper in PAPERS:
            predicted = build_term_texts(models[condition][paper], mode=mode)
            expert_scores = {
                expert: score_terms(
                    build_term_texts(human[paper][expert], mode=mode),
                    predicted,
                    cache,
                )
                for expert in EXPERTS
            }
            paper_mean = mean_exact(list(expert_scores.values()), expected=4)
            paper_results[paper] = {
                "predicted_count": len(predicted),
                "expert_scores": expert_scores,
                "mean": paper_mean,
            }
            paper_scores.append(paper_mean)
        modes[mode] = {
            "macro_mean": mean_exact(paper_scores, expected=7),
            "paper_scores": paper_results,
        }
    return {"key": condition, "label": CONDITION_LABELS[condition], "coverage": "7/7", "modes": modes}


def evaluate_human_agreement(
    human: Mapping[str, Mapping[str, Sequence[Mapping[str, Any]]]],
    cache: SharedEmbeddingCache,
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for mode in MODES:
        papers: dict[str, Any] = {}
        paper_means: list[float] = []
        for paper in PAPERS:
            pair_scores = {
                f"{left} vs {right}": score_terms(
                    build_term_texts(human[paper][left], mode=mode),
                    build_term_texts(human[paper][right], mode=mode),
                    cache,
                )
                for left, right in combinations(EXPERTS, 2)
            }
            paper_mean = mean_exact(list(pair_scores.values()), expected=6)
            papers[paper] = {"pair_scores": pair_scores, "mean": paper_mean}
            paper_means.append(paper_mean)
        result[mode] = {"macro_mean": mean_exact(paper_means, expected=7), "paper_scores": papers}
    return result


def write_csv(path: Path, conditions: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["condition", "label", "coverage", "label_only_SBMS_F1", "label_plus_evidence_SBMS_F1"])
        for condition in conditions:
            writer.writerow(
                [
                    condition["key"],
                    condition["label"],
                    condition["coverage"],
                    f"{condition['modes']['label_only']['macro_mean']:.10f}",
                    f"{condition['modes']['label_plus_evidence']['macro_mean']:.10f}",
                ]
            )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-root", type=Path, required=True)
    parser.add_argument("--human-workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--llm-config", type=Path, default=REPO_ROOT / "llm.json")
    parser.add_argument("--gemini-run-root", type=Path)
    args = parser.parse_args()
    run_root = args.run_root.expanduser().resolve()
    gemini_run_root = args.gemini_run_root.expanduser().resolve() if args.gemini_run_root else None
    human_workbook = args.human_workbook.expanduser().resolve()
    human = load_human_standard(human_workbook)
    paths = resolve_condition_paths(run_root, gemini_run_root=gemini_run_root)
    condition_keys = tuple(paths)
    models = load_condition_models(paths)
    all_texts = collect_texts(human, models, conditions=condition_keys)
    unique_texts = list(dict.fromkeys(all_texts))
    print(
        f"Embedding {len(unique_texts)} unique texts once for all {len(condition_keys)} conditions",
        flush=True,
    )
    cache = SharedEmbeddingCache(args.llm_config.expanduser().resolve())
    cache.prepare(unique_texts)
    conditions = [evaluate_condition(condition, human, models, cache) for condition in condition_keys]
    payload = {
        "generated_at": utc_now(),
        "protocol": {
            "run_root": str(run_root),
            "human_workbook": str(human_workbook),
            "experts": list(EXPERTS),
            "papers": list(PAPERS),
            "embedding_provider": cache.config.provider_name,
            "embedding_model": cache.config.model_name,
            "unique_text_count": len(unique_texts),
            "embedding_dimension": cache.dimension,
            "shared_cache": "each unique text embedded exactly once in this evaluation run",
            "text_modes": {
                "label_only": "cell_subtype only",
                "label_plus_evidence": "cell_subtype + newline + 'Evidence: ' + whitespace-normalized evidence",
            },
            "paper_aggregation": "mean of model-vs-four-expert SBMS F1",
            "overall_aggregation": "macro mean over seven papers",
            "pair_count": len(condition_keys) * 7 * 2 * 4,
            "gemini_run_root": str(gemini_run_root) if gemini_run_root else None,
            "condition_output_paths": {
                condition: {paper: str(path) for paper, path in paper_paths.items()}
                for condition, paper_paths in paths.items()
            },
        },
        "human_agreement": evaluate_human_agreement(human, cache),
        "conditions": conditions,
    }
    result_json, result_csv = resolve_result_paths(
        args.output_dir.expanduser().resolve(),
        include_gemini=gemini_run_root is not None,
    )
    write_json(result_json, payload)
    write_csv(result_csv, conditions)
    print(json.dumps({"result_json": str(result_json), "result_csv": str(result_csv)}, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
